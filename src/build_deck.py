"""Presentación del caso: PPTX editable (texto nativo) + PDF idéntico exportado con LibreOffice.

Uso:  .venv\\Scripts\\python.exe src\\build_deck.py

Salidas:
  deck/presentacion.pptx y deck/presentacion.pdf           (16:9, relato completo)
  deck/carrusel_linkedin.pptx y deck/carrusel_linkedin.pdf (4:5, versión corta para LinkedIn)
Las cifras salen de kpis.json con la regla única de formato.
"""
import json
import os
import shutil
import subprocess
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from build_docs import AUTOR, limpiar_metadatos, protagonista, resultados
from config import FECHA_CORTE, FIGURES, PROCESSED, ROOT

DECK = ROOT / "deck"
# LibreOffice: variable de entorno SOFFICE, instalación estándar o el ejecutable disponible en el PATH
SOFFICE = Path(os.environ.get("SOFFICE") or (r"C:\Program Files\LibreOffice\program\soffice.exe"
                                            if Path(r"C:\Program Files\LibreOffice\program\soffice.exe").exists()
                                            else shutil.which("soffice") or "soffice"))
FUENTE = "Segoe UI"
BG, INK, INK2, MUTED, LINE = "FCFCFB", "0B0B0B", "52514E", "898781", "E1E0D9"
BLUE, BLUE_SOFT, BLUE_DARK, RED, GRAY = "2A78D6", "EAF2FC", "184F95", "E34948", "C3C2B7"
RAMPA = ["86B6EF", "5598E7", "3987E5", "256ABF", "184F95"]


def rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


APP_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           '<Slides>{slides}</Slides><Company></Company></Properties>')


class Deck:
    def __init__(self, ancho_in: float, alto_in: float):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(ancho_in), Inches(alto_in)
        self.W, self.H = ancho_in, alto_in
        self.n = 0
        props = self.prs.core_properties
        props.author, props.last_modified_by, props.title = AUTOR, AUTOR, "¿Quién puede pagar hoy un crédito hipotecario UVA?"
        props.subject, props.keywords, props.comments = "Caso de análisis de datos", "crédito hipotecario, UVA, EPH, vivienda", ""
        # Fechas del caso, no las de la plantilla base
        props.created = props.modified = datetime.combine(FECHA_CORTE, datetime.min.time())
        props.revision = 1

    # --- piezas -----------------------------------------------------------------------------------------
    def slide(self, kicker: str | None = None, titulo: str | None = None, fuente: str | None = None, notas: str = ""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(BG)
        self.n += 1
        m = 0.6
        if kicker:
            self.texto(s, m, 0.42, self.W - 2 * m, 0.35, [kicker.upper()], size=11, color=BLUE, bold=True)
        if titulo:
            self.texto(s, m, 0.72, self.W - 2 * m, 1.1, [titulo], size=28 if self.W > 10 else 26, bold=True, color=INK)
        if fuente:
            self.texto(s, m, self.H - 0.5, self.W - 2 * m - 0.8, 0.3, ["Fuente: " + fuente], size=9, color=MUTED)
        self.texto(s, self.W - m - 0.6, self.H - 0.5, 0.6, 0.3, [str(self.n)], size=9, color=MUTED, align=PP_ALIGN.RIGHT)
        if notas:
            s.notes_slide.notes_text_frame.text = notas
        return s

    def texto(self, s, x, y, w, h, parrafos, size=16, color=INK, bold=False, align=None, anchor=MSO_ANCHOR.TOP, interlineado=1.1):
        caja = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = caja.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, attr, 0)
        for i, par in enumerate(parrafos):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = interlineado
            if align is not None:
                p.alignment = align
            runs = par if isinstance(par, list) else [(par, {})]
            for contenido, estilo in runs:
                r = p.add_run()
                r.text = contenido
                f = r.font
                f.name = FUENTE
                f.size = Pt(estilo.get("size", size))
                f.bold = estilo.get("bold", bold)
                f.color.rgb = rgb(estilo.get("color", color))
            if isinstance(par, list) and par and par[0][1].get("espacio"):
                p.space_before = Pt(par[0][1]["espacio"])
        return caja

    def caja(self, s, x, y, w, h, relleno=None, borde=LINE, radio=True):
        forma = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radio else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if radio:
            forma.adjustments[0] = 0.06
        forma.shadow.inherit = False
        for ref in forma._element.xpath("./p:style/a:effectRef"):
            ref.set("idx", "0")
        if relleno:
            forma.fill.solid()
            forma.fill.fore_color.rgb = rgb(relleno)
        else:
            forma.fill.background()
        if borde:
            forma.line.color.rgb = rgb(borde)
            forma.line.width = Pt(1)
        else:
            forma.line.fill.background()
        return forma

    def tarjeta(self, s, x, y, w, h, valor, etiqueta, detalle="", color_valor=INK, relleno="FFFFFF", size_valor=30):
        self.caja(s, x, y, w, h, relleno=relleno)
        self.texto(s, x + 0.22, y + 0.2, w - 0.44, 0.7, [valor], size=size_valor, bold=True, color=color_valor)
        self.texto(s, x + 0.22, y + 0.2 + size_valor / 55, w - 0.44, h - 0.4 - size_valor / 55,
                   [etiqueta] + ([[(detalle, {"size": 11, "color": MUTED, "espacio": 4})]] if detalle else []), size=13, color=INK2)

    def figura(self, s, archivo: str, x, y, w=None, h=None):
        ruta = FIGURES / f"{archivo}.png"
        kwargs = {}
        if w:
            kwargs["width"] = Inches(w)
        if h:
            kwargs["height"] = Inches(h)
        pic = s.shapes.add_picture(str(ruta), Inches(x), Inches(y), **kwargs)
        if w and h is None and pic.height > Inches(self.H - y - 0.6):   # si no entra, ajustar por alto
            ratio = pic.width / pic.height
            pic.height = Inches(self.H - y - 0.6)
            pic.width = Emu(int(pic.height * ratio))
        return pic

    def barras(self, s, x, y, w, h, categorias, valores, etiquetas, colores, max_valor=None):
        datos = CategoryChartData()
        datos.categories = categorias
        datos.add_series("Hogares", valores)
        grafico = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), datos).chart
        grafico.has_legend = False
        grafico.has_title = False
        grafico.font.name, grafico.font.size, grafico.font.color.rgb = FUENTE, Pt(13), rgb(INK2)
        eje_c, eje_v = grafico.category_axis, grafico.value_axis
        eje_c.reverse_order = True
        eje_c.format.line.color.rgb = rgb(LINE)
        eje_c.has_major_gridlines = False
        eje_c.tick_labels.font.size = Pt(13)
        eje_v.visible = False
        eje_v.has_major_gridlines = False
        eje_v.minimum_scale = 0
        if max_valor:
            eje_v.maximum_scale = max_valor
        plot = grafico.plots[0]
        plot.gap_width = 45
        serie = plot.series[0]
        for i, (color, etiqueta) in enumerate(zip(colores, etiquetas)):
            punto = serie.points[i]
            punto.format.fill.solid()
            punto.format.fill.fore_color.rgb = rgb(color)
            punto.format.line.fill.background()
            dl = punto.data_label
            dl.has_text_frame = True
            dl.text_frame.text = etiqueta
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            for p in dl.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.name, r.font.color.rgb, r.font.bold = Pt(13), FUENTE, rgb(INK2), False
        return grafico

    def guardar(self, pptx: Path):
        pptx.parent.mkdir(parents=True, exist_ok=True)
        # Propiedades extendidas mínimas: la plantilla base trae nombre de aplicación y formato que no corresponden
        for parte in self.prs.part.package.iter_parts():
            if str(parte.partname) == "/docProps/app.xml":
                parte._blob = APP_XML.format(slides=self.n).encode("utf-8")
        self.prs.save(str(pptx))


# ---------------------------------------------------------------------------
# Presentación 16:9
# ---------------------------------------------------------------------------
def presentacion(r: dict, prot: dict, K: dict) -> Deck:
    d = Deck(13.333, 7.5)
    fecha = FECHA_CORTE.strftime("%d-%m-%Y")
    emb = K["base"]["embudo"]

    # 1 · Portada
    s = d.slide(notas="Caso de análisis con datos públicos para un banco ficticio que relanza su línea hipotecaria UVA.")
    d.caja(s, 0, 0, 0.18, 7.5, relleno=BLUE, borde=None, radio=False)
    d.texto(s, 0.9, 1.3, 11, 0.4, ["CASO DE ANÁLISIS · BANCO [EMPRESA] (FICTICIO)"], size=13, color=BLUE, bold=True)
    d.texto(s, 0.9, 1.85, 11.2, 2.2, ["¿Quién puede pagar hoy un crédito hipotecario UVA?"], size=48, bold=True, interlineado=1.0)
    d.texto(s, 0.9, 4.15, 10.5, 1.0, ["Hogares inquilinos de CABA, Partidos del GBA y Gran Córdoba: cuántos califican, qué los frena y qué producto amplía el mercado."], size=19, color=INK2)
    d.texto(s, 0.9, 6.3, 11, 0.6, [[(AUTOR, {"bold": True, "color": INK}), (f"   ·   Datos al {fecha}   ·   INDEC · BCRA · datos.gob.ar · Zonaprop Index", {"color": MUTED})]], size=13)

    # 2 · Gancho
    s = d.slide(notas="El número que abre la conversación: la gran mayoría de los inquilinos no califica con las condiciones de hoy.",
                fuente="INDEC (EPH 2T-2025 a 1T-2026), BCRA, Zonaprop Index; elaboración propia.")
    d.texto(s, 0.9, 1.3, 6.5, 1.5, [f"1 de cada {r['uno_de_cada']}"], size=72, bold=True, color=BLUE, interlineado=0.95)
    d.texto(s, 0.9, 2.95, 6.0, 1.4, ["hogares que alquilan podría tomar hoy un crédito hipotecario UVA."], size=26, color=INK, interlineado=1.05)
    d.texto(s, 0.9, 4.6, 5.8, 1.6, [f"Son {r['elegibles']} de {r['inquilinos']} de hogares inquilinos. El resto no llega a la cuota, al anticipo o no puede demostrar su ingreso."], size=16, color=INK2)
    d.tarjeta(s, 7.6, 1.3, 4.9, 1.55, r["inquilinos"], "hogares inquilinos en los 3 mercados", "base de todos los porcentajes")
    d.tarjeta(s, 7.6, 3.05, 4.9, 1.55, r["cuota_pct"], "puede pagar la cuota", "cuota hasta 25% del ingreso")
    d.tarjeta(s, 7.6, 4.8, 4.9, 1.55, r["fgs"], "de fondeo del programa oficial", f"alcanza para ~{r['fgs_creditos']} créditos")

    # 3 · Protagonista
    s = d.slide("Una historia", "Carla y Diego ganan lo típico y no llegan", fuente="Hogar ilustrativo construido con medianas de la EPH y precios del Zonaprop Index (GBA Oeste).",
                notas="Hogar ilustrativo: pareja con hijos con el ingreso mediano de su grupo, buscando 60 m² en el GBA Oeste.")
    d.texto(s, 0.6, 1.55, 12.1, 0.9, [f"Pareja con un hijo, alquilan en el oeste del conurbano y ganan {prot['ingreso']} por mes entre los dos. Buscan un 3 ambientes de {prot['m2']} m² que cuesta {prot['base']['precio_usd']}."], size=18, color=INK2)
    b, pp, pb = prot["base"], prot["propuesto"], prot["propuesto_barata"]
    cols = [
        ("Hoy (20 años, 75%)", f"{b['ratio']}", "del ingreso se iría en la cuota", f"Cuota {b['cuota']} · efectivo {b['efectivo']} ({b['anios_ahorro']} años ahorrando el 20%)", RED),
        ("Con 30 años y 80%", f"{pp['ratio']}", "del ingreso: todavía no alcanza", f"Cuota {pp['cuota']} · efectivo {pp['efectivo']}", RED),
        ("Y una vivienda 20% más barata", f"{pb['ratio']}", "del ingreso: " + ("califican" if pb["pasa_cuota"] and pb["pasa_anticipo"] else "siguen sin calificar"), f"Efectivo equivalente a {pb['meses']} meses de ingreso", BLUE),
    ]
    for i, (t, v, l, det, color) in enumerate(cols):
        x = 0.6 + i * 4.1
        d.caja(s, x, 2.75, 3.85, 3.6, relleno="FFFFFF")
        d.texto(s, x + 0.25, 2.95, 3.4, 0.4, [t], size=14, bold=True, color=INK)
        d.texto(s, x + 0.25, 3.45, 3.4, 1.1, [v], size=54, bold=True, color=color)
        d.texto(s, x + 0.25, 4.55, 3.4, 0.5, [l], size=15, color=INK)
        d.texto(s, x + 0.25, 5.2, 3.4, 1.0, [det], size=12, color=MUTED)

    # 4 · La pregunta
    s = d.slide("La pregunta del banco", "Hay fondeo barato y 120 días para colocarlo: ¿a quién y con qué producto?",
                fuente="Referencias de supuestos: condiciones publicadas del programa de fondeo y de Banco Nación (detalle en la documentación del repositorio).",
                notas="Tres preguntas de negocio que ordenan todo el análisis.")
    preguntas = [("1", "¿Cuántos hogares inquilinos podrían pagar hoy un crédito UVA y dónde están?"),
                 ("2", "¿Qué los frena: la cuota, el anticipo, la formalidad o la edad?"),
                 ("3", "¿Qué producto y qué política de riesgo amplían el mercado sin aumentar la mora esperada?")]
    for i, (num, txt) in enumerate(preguntas):
        y = 2.0 + i * 1.15
        d.caja(s, 0.6, y, 0.8, 0.8, relleno=BLUE_SOFT, borde=None)
        d.texto(s, 0.6, y + 0.12, 0.8, 0.6, [num], size=26, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
        d.texto(s, 1.65, y + 0.12, 6.4, 0.9, [txt], size=18, color=INK)
    d.caja(s, 8.4, 2.0, 4.3, 3.1, relleno="FFFFFF")
    d.texto(s, 8.65, 2.2, 3.9, 2.8, [[("Escenario base", {"bold": True, "color": INK})], "UVA + 7,5% · 20 años · 75% financiado", "Cuota hasta 25% del ingreso",
                                     "Anticipo + gastos hasta 12 ingresos", "Recibo de sueldo (jefe/a o cónyuge)", "Edad al terminar hasta 85 años"], size=14, color=INK2, interlineado=1.3)
    d.texto(s, 0.6, 5.65, 12, 0.8, ["Alcance: CABA, Partidos del GBA (abierto en corredores Norte, Oeste y Sur) y Gran Córdoba."], size=14, color=MUTED)

    # 5 · Los datos
    s = d.slide("Los datos", "Datos públicos, limpiados con reglas escritas y validados contra cifras oficiales",
                fuente="INDEC (EPH, Censo 2022), BCRA (API de Estadísticas), datos.gob.ar (salarios, SIPA, dólar BNA), GCBA, Zonaprop Index.",
                notas="Antes de analizar, las cifras oficiales del INDEC se reprodujeron exactamente desde los microdatos.")
    fuentes = [("EPH · INDEC", "ingresos, tenencia, trabajo y aportes de los hogares"), ("Censo 2022 · INDEC", "validación del universo y peso de cada partido"),
               ("BCRA", "UVA, tasas, stock y destino del crédito"), ("datos.gob.ar", "índice de salarios, SIPA, dólar Banco Nación"),
               ("Zonaprop Index", "precio del m² en CABA, 31 partidos y Córdoba")]
    for i, (f, det) in enumerate(fuentes):
        d.texto(s, 0.6, 2.0 + i * 0.72, 6.4, 0.7, [[(f, {"bold": True, "color": INK}), ("  ·  " + det, {"color": INK2})]], size=14)
    controles = [("4 de 4", "cifras oficiales del INDEC reproducidas exactamente"), ("2 a 4 pp", "menos inquilinos que el Censo: cantidades conservadoras"),
                 (r["obs_inq"], "observaciones de hogares inquilinos con ingreso")]
    for i, (v, l) in enumerate(controles):
        d.tarjeta(s, 7.2, 1.85 + i * 1.5, 5.5, 1.3, v, l, size_valor=26)

    # 6 · La respuesta: embudo
    s = d.slide("La respuesta", f"Solo {r['elegibles']} hogares pasan los cuatro filtros",
                fuente="INDEC (EPH), BCRA, Zonaprop Index; elaboración propia. Escenario base.",
                notas="La mayor caída está en la cuota; después pesan el anticipo y la falta de recibo de sueldo.")
    etapas = ["Hogares inquilinos", "Pagan la cuota", "Juntan el anticipo", "Ingreso demostrable", "Cumplen la edad"]
    valores = [e["hogares"] for e in emb]
    from formato import fmt_cantidad, fmt_pct
    etiquetas = [f"{fmt_cantidad(e['hogares'])} · {fmt_pct(e['pct_del_total'])}" for e in emb]
    d.barras(s, 0.6, 1.8, 8.2, 4.6, etapas, valores, etiquetas, RAMPA, max_valor=valores[0] * 1.35)
    d.caja(s, 9.2, 2.0, 3.5, 4.2, relleno=BLUE_SOFT, borde=None)
    d.texto(s, 9.45, 2.25, 3.05, 3.8, [[(f"{r['no_cuota']} de cada 10", {"bold": True, "size": 24, "color": BLUE_DARK})], "no llegan a la cuota.",
                                      [(f"{r['sin_formal_pct']}", {"bold": True, "size": 24, "color": BLUE_DARK, "espacio": 14})], "de los que pagan y juntan el anticipo no tiene recibo de sueldo.",
                                      [(f"IC {r['ic_inf']}–{r['ic_sup']}", {"bold": True, "size": 18, "color": BLUE_DARK, "espacio": 14})], "margen de error del resultado"], size=14, color=INK2)

    # 7 · Por qué
    s = d.slide("Por qué tan pocos", "Lo que frena es juntar el efectivo inicial frente al precio",
                fuente="INDEC (EPH), BCRA, Zonaprop Index; elaboración propia.",
                notas="Cuota y anticipo dependen del mismo cociente precio/ingreso; el anticipo es el límite más exigente.")
    d.tarjeta(s, 0.6, 1.85, 3.9, 1.9, r["mediana_precio_ing"], "ingresos mensuales cuesta la vivienda del hogar típico", size_valor=48)
    d.tarjeta(s, 0.6, 3.95, 3.9, 1.9, r["tope_anticipo"], "ingresos es lo máximo que permite el anticipo (y 41,4 la cuota)", size_valor=48, color_valor=BLUE)
    d.figura(s, "03_precio_en_ingresos", 4.8, 1.8, w=7.9)

    # 8 · Superficie
    s = d.slide("La vivienda", f"Con 45 m² calificaría el {r['calor']['Total_45']}; con 75 m², el {r['calor']['Total_75']}",
                fuente="INDEC (EPH), Zonaprop Index, BCRA; elaboración propia.",
                notas="En lugar de fijar una regla de m², se muestra la elegibilidad para cada superficie.")
    d.figura(s, "03_calor_m2_zona", 0.6, 1.7, w=12.1)

    # 9 · Qué pasa si
    s = d.slide("Qué destraba", "Lo que suma hogares actúa sobre el anticipo o el precio",
                fuente="INDEC (EPH), BCRA, Zonaprop Index; elaboración propia.",
                notas="Cada medida se prueba de a una frente al escenario base.")
    d.figura(s, "04_que_pasa_si", 1.4, 1.7, w=10.5)

    # 10 · Producto
    s = d.slide("El producto", f"30 años y 80% financiado: de {r['elegibles']} a {r['prop']} hogares sin subir el tope de cuota",
                fuente="INDEC (EPH), BCRA, Zonaprop Index; elaboración propia. Tope de cuota 25% en los tres escenarios.",
                notas="La grilla de plazo por % financiado muestra que 30 años con 80% es la mejor combinación con cuota hasta 25%.")
    pk = K["producto_propuesto"]
    d.barras(s, 0.6, 1.9, 7.4, 3.6, ["Hoy: 20 años, 75%", "Propuesto: 30 años, 80%", "Propuesto + independientes"],
             [K["base"]["elegibles_hogares"], pk["elegibles_hogares"], pk["con_independientes_hogares"]],
             [r["elegibles"], f"{r['prop']} ({r['prop_delta']})", f"{r['prop_indep']} ({r['prop_indep_delta']})"], [GRAY, BLUE, BLUE],
             max_valor=pk["con_independientes_hogares"] * 1.35)
    d.figura(s, "04_grilla_plazo_ltv", 8.35, 1.95, w=4.4)

    # 11 · A quién
    s = d.slide("A quién", "Dos ingresos hacen la diferencia; la zona pesa menos", fuente="INDEC (EPH); elaboración propia. Asociaciones descriptivas, no causas.",
                notas="Mirada descriptiva y neutral por tipo de hogar y trabajos; independientes en un módulo aparte.")
    tarjetas = [
        (r["trab"]["2 o más personas trabajando"]["pct"], "elegible donde trabajan dos personas o más", f"frente a {r['trab']['1 persona con 1 trabajo']['pct']} con una persona y un trabajo"),
        (r["perfil"]["Pareja sin hijos|Todas"]["pct"], "elegible en parejas sin hijos", f"hogares de un solo adulto: {r['perfil']['Unipersonal|Todas']['pct']} y {r['perfil']['Monoparental|Todas']['pct']}"),
        (r["ind"]["pagan"], "hogares de independientes pagarían cuota y anticipo", "pero no tienen recibo de sueldo"),
        (r["zonas"]["GBA Oeste"]["pct"], "elegible en GBA Oeste, la zona más alta", f"GBA Norte {r['zonas']['GBA Norte']['pct']}; diferencia no concluyente"),
    ]
    for i, (v, l, det) in enumerate(tarjetas):
        d.tarjeta(s, 0.6 + (i % 2) * 6.15, 1.9 + (i // 2) * 2.3, 5.95, 2.05, v, l, det, size_valor=40, color_valor=BLUE if i < 3 else INK)

    # 12 · Riesgo
    s = d.slide("El riesgo", "Si la UVA le gana al sueldo, la cuota pesa más",
                fuente="BCRA (UVA), INDEC (índice de salarios registrados); elaboración propia.",
                notas="Por eso el producto mantiene el tope de cuota en 25%: con 30% al otorgar, en un episodio así la cuota llegaría cerca del 40%.")
    d.figura(s, "03_stress_uva_salarios", 0.6, 1.8, w=8.0)
    d.caja(s, 8.95, 2.0, 3.75, 4.0, relleno=BLUE_SOFT, borde=None)
    d.texto(s, 9.2, 2.25, 3.3, 3.6, [[(r["stress24"], {"bold": True, "size": 34, "color": BLUE_DARK})], "del ingreso a los 24 meses para un crédito de sep-2023 que empezó en 25%.",
                                     [(r["stress24_sobre30"], {"bold": True, "size": 34, "color": BLUE_DARK, "espacio": 16})], "de los períodos de 24 meses desde 2016 terminó por encima del 30%."], size=14, color=INK2)

    # 13 · Recomendación
    s = d.slide("Recomendación", "Recomiendo 30 años, 80% financiado y cuota hasta 25%, con foco en dos ingresos y una línea para independientes",
                fuente="Cifras: kpis.json del proyecto; elaboración propia.", notas="Recomiendo este producto porque suma hogares sin subir el tope de cuota, baja el efectivo inicial y protege frente al desacople.")
    porques = [("Suma más hogares", f"De {r['elegibles']} a {r['prop']}, y a {r['prop_indep']} con independientes, sin subir el tope de cuota."),
               ("Baja el efectivo inicial", "Financiar el 80% a 30 años reduce el anticipo y mantiene la cuota dentro del tope."),
               ("Cuida el riesgo", "Tope de 25% y opción de extender el plazo si la cuota supera el 30% del sueldo."),
               ("Permite seleccionar", f"El fondeo alcanza para ~{r['fgs_creditos']} créditos frente a {r['prop']} hogares elegibles.")]
    for i, (t, det) in enumerate(porques):
        x = 0.6 + (i % 2) * 6.15
        y = 2.3 + (i // 2) * 2.05
        d.caja(s, x, y, 5.95, 1.8, relleno="FFFFFF")
        d.texto(s, x + 0.25, y + 0.2, 5.45, 0.45, [t], size=18, bold=True, color=BLUE)
        d.texto(s, x + 0.25, y + 0.7, 5.45, 1.0, [det], size=14, color=INK2)

    # 14 · Próximos pasos y límites
    s = d.slide("Próximos pasos", "Qué haría en los 120 días del programa y qué no dice este análisis",
                fuente="Límites detallados en la documentación del repositorio.", notas="Plan de acción con responsables e indicadores, y límites del estudio.")
    pasos = ["Semanas 1-2 · aprobar producto y regla de sumar ingresos de la pareja", "Semanas 1-4 · precalificación digital con el simulador",
             "Semanas 3-6 · línea de independientes con 12-24 meses de facturación o aportes", "Semanas 3-8 · campañas: dos ingresos; GBA Oeste y Córdoba por proporción, CABA por volumen",
             "Semanas 6-12 · ahorro programado en UVA para casi elegibles", "Todo el período · tablero de colocación y mora temprana por línea"]
    d.texto(s, 0.6, 1.9, 6.6, 4.5, [[("Plan de acción", {"bold": True, "color": INK, "size": 17})]] + ["•  " + p for p in pasos], size=14, color=INK2, interlineado=1.35)
    limites = ["La encuesta no mide ahorros, deudas ni historial crediticio", "Precios de publicación, no de escritura",
               f"Cantidades conservadoras: entre {r['rango_censo']} ajustado al Censo", "En el GBA se asume igual distribución de ingresos por partido", "Válido a la fecha de corte; tasas y precios cambian"]
    d.caja(s, 7.6, 1.9, 5.1, 4.4, relleno="F0EFEC", borde=None)
    d.texto(s, 7.85, 2.1, 4.6, 4.1, [[("Límites", {"bold": True, "color": INK, "size": 17})]] + ["•  " + x for x in limites], size=13.5, color=INK2, interlineado=1.35)

    # 15 · Cómo se hizo
    s = d.slide("Cómo se hizo", "Decisiones humanas en cada etapa, cálculos automatizados y controles verificables",
                fuente="Proceso completo documentado en el repositorio público del caso.", notas="Participación del autor en cuatro puntos de control.")
    bloques = [("Qué decidió el autor", "Tema, cliente y preguntas. Alcance: sacar Gran Rosario y abrir el GBA por corredor. Escenario base. Qué hipótesis mantener, reformular o pasar a contexto. Qué controles sumar."),
               ("Qué se automatizó", "Descarga de fuentes públicas con huella digital, limpieza en SQL, motor de elegibilidad único, gráficos, documentos y dashboard."),
               ("Cómo se verificó", f"Cifras del INDEC reproducidas exactamente, comparación con el Censo 2022, imputación de ingresos no declarados ({r['nr']['pct']}), corrección por subdeclaración ({r['sub']['mediana']}–{r['sub']['promedio']}) y test de paridad del simulador.")]
    for i, (t, det) in enumerate(bloques):
        x = 0.6 + i * 4.1
        d.caja(s, x, 1.95, 3.85, 4.3, relleno="FFFFFF")
        d.texto(s, x + 0.25, 2.15, 3.35, 0.45, [t], size=17, bold=True, color=INK)
        d.texto(s, x + 0.25, 2.7, 3.35, 3.4, [det], size=14, color=INK2, interlineado=1.25)

    # 16 · Fuentes y cierre
    s = d.slide("Fuentes de datos y material", "Todo el caso se puede recorrer y repetir", notas="Links al dashboard, documentos y repositorio.")
    fuentes = ["INDEC · Encuesta Permanente de Hogares (2T-2025 a 1T-2026) y Censo 2022", "BCRA · API de Estadísticas v4.0 (UVA, tasas, stock y destino del crédito, inflación, dólar)",
               "datos.gob.ar · Índice de Salarios, RIPTE, remuneraciones del SIPA y dólar Banco Nación", "GCBA · Departamentos en venta 2001-2020", "Zonaprop Index · reportes CABA, GBA y Córdoba 2026"]
    d.texto(s, 0.6, 1.9, 7.2, 3.6, ["•  " + f for f in fuentes], size=14.5, color=INK2, interlineado=1.4)
    d.caja(s, 8.2, 1.9, 4.5, 3.3, relleno=BLUE_SOFT, borde=None)
    d.texto(s, 8.45, 2.1, 4.0, 3.0, [[("Material del caso", {"bold": True, "color": INK, "size": 16})], "Dashboard interactivo con simulador", "Resumen en simple",
                                     "Código, documentación y fuentes en GitHub"], size=14, color=INK2, interlineado=1.45)
    d.texto(s, 0.6, 5.7, 12.1, 1.0, ["Ejercicio analítico independiente con datos públicos, para portfolio profesional. Banco [EMPRESA] es ficticio y el hogar de la historia es ilustrativo. No es asesoramiento financiero."], size=12, color=MUTED)
    d.texto(s, 0.6, 6.55, 12.1, 0.4, [AUTOR], size=14, bold=True, color=INK)
    return d


# ---------------------------------------------------------------------------
# Carrusel LinkedIn 4:5
# ---------------------------------------------------------------------------
def carrusel(r: dict, prot: dict, K: dict) -> Deck:
    d = Deck(7.5, 9.375)
    m, w = 0.6, 7.5 - 1.2

    def pie(s, texto_fuente):
        d.texto(s, m, 8.72, w - 0.7, 0.4, [texto_fuente], size=9, color=MUTED)

    s = d.slide()
    d.caja(s, 0, 0, 7.5, 0.16, relleno=BLUE, borde=None, radio=False)
    d.texto(s, m, 1.1, w, 0.4, ["CASO DE ANÁLISIS DE DATOS"], size=13, color=BLUE, bold=True)
    d.texto(s, m, 1.6, w, 3.2, ["¿Quién puede pagar hoy un crédito hipotecario UVA?"], size=44, bold=True, interlineado=1.0)
    d.texto(s, m, 4.9, w, 1.6, ["Hogares que alquilan en CABA, GBA y Gran Córdoba. Datos públicos del INDEC, BCRA y Zonaprop."], size=19, color=INK2)
    d.texto(s, m, 8.0, w, 0.5, [[(AUTOR, {"bold": True}), ("   ·   deslizá →", {"color": MUTED})]], size=14)

    s = d.slide()
    d.texto(s, m, 1.4, w, 1.5, [f"1 de cada {r['uno_de_cada']}"], size=70, bold=True, color=BLUE)
    d.texto(s, m, 3.3, w, 1.8, ["hogares que alquilan podría tomar hoy un crédito hipotecario UVA."], size=28, interlineado=1.05)
    d.texto(s, m, 5.4, w, 1.5, [f"{r['elegibles']} de {r['inquilinos']} de hogares inquilinos."], size=20, color=INK2)
    pie(s, "Fuente: INDEC (EPH), BCRA, Zonaprop Index; elaboración propia.")

    s = d.slide()
    d.texto(s, m, 0.8, w, 1.2, ["Carla y Diego ganan lo típico y no llegan"], size=32, bold=True)
    b, pp, pb = prot["base"], prot["propuesto"], prot["propuesto_barata"]
    filas = [(b["ratio"], "del ingreso iría a la cuota hoy (el banco acepta 25%)", RED), (b["efectivo"], f"de anticipo y gastos: {b['anios_ahorro']} años ahorrando el 20%", RED),
             (pp["ratio"], "con 30 años y 80% financiado: todavía no", RED), (pb["ratio"], "con una vivienda 20% más barata: " + ("califican" if pb["pasa_cuota"] and pb["pasa_anticipo"] else "tampoco"), BLUE)]
    for i, (v, l, c) in enumerate(filas):
        y = 2.2 + i * 1.55
        d.texto(s, m, y, 3.0, 1.0, [v], size=28 if len(v) > 5 else 34, bold=True, color=c)
        d.texto(s, m + 3.1, y + 0.08, w - 3.1, 1.2, [l], size=16, color=INK2)
    pie(s, f"Hogar ilustrativo: pareja con hijos, ingreso {prot['ingreso']}, 60 m² en GBA Oeste.")

    s = d.slide()
    d.texto(s, m, 0.8, w, 1.6, ["Lo que frena es el efectivo inicial"], size=32, bold=True)
    q = {x["pregunta"]: x["diferencia"] for x in K["que_pasa_si"]}
    medidas = [("Vivienda 20% más barata", "¿Y si se compra una vivienda 20% más barata (usada o más chica)?"),
               ("Aceptar independientes", "¿Y si acepta monotributistas y autónomos con aportes?"),
               ("30 años y 80% financiado", "¿Y si hace las dos cosas: 30 años y 80%?"),
               ("Ahorro para el anticipo", "¿Y si un plan de ahorro permite juntar 24 meses de ingreso?"),
               ("Bajar la tasa a 5,5%", "¿Y si la tasa baja de 7,5% a 5,5%?")]
    valores = [max(q[k], 0) for _, k in medidas]
    from formato import fmt_cantidad
    etiquetas = [("+" + fmt_cantidad(v) + " hogares") if v > 0.5 else "sin cambio" for v in valores]
    d.barras(s, m, 2.5, w, 4.8, [e for e, _ in medidas], valores, etiquetas, [BLUE, BLUE, BLUE, BLUE, GRAY], max_valor=max(valores) * 1.6)
    d.texto(s, m, 7.45, w, 1.0, ["Hogares elegibles que se suman frente a hoy, cambiando una condición por vez."], size=15, color=INK2)
    pie(s, "Fuente: INDEC (EPH), BCRA, Zonaprop Index; elaboración propia.")

    s = d.slide()
    d.texto(s, m, 0.8, w, 1.6, ["Un producto mejor diseñado suma hogares sin subir la cuota"], size=30, bold=True)
    pk = K["producto_propuesto"]
    d.barras(s, m, 2.7, w, 3.8, ["Hoy: 20 años, 75%", "30 años, 80%", "+ independientes"],
             [K["base"]["elegibles_hogares"], pk["elegibles_hogares"], pk["con_independientes_hogares"]],
             [r["elegibles"], r["prop"], r["prop_indep"]], [GRAY, BLUE, BLUE], max_valor=pk["con_independientes_hogares"] * 1.4)
    d.texto(s, m, 6.8, w, 1.4, ["Monotributistas y autónomos con aportes, en una línea separada con su propia política de riesgo."], size=17, color=INK2)
    pie(s, "Fuente: INDEC (EPH), BCRA, Zonaprop Index; elaboración propia. Tope de cuota 25%.")

    s = d.slide()
    d.texto(s, m, 0.8, w, 1.6, ["Si la UVA le gana a los sueldos, la cuota pesa más"], size=30, bold=True)
    d.texto(s, m, 2.6, w, 1.2, [[("25%", {"color": INK2}), ("  →  ", {"color": MUTED}), (r["stress24"], {"color": RED})]], size=60, bold=True)
    d.texto(s, m, 3.9, w, 1.4, ["del ingreso para la cuota, en dos años, en un crédito tomado en septiembre de 2023."], size=20, color=INK2)
    d.texto(s, m, 5.5, w, 2.4, [f"El {r['stress24_sobre30']} de los períodos de 24 meses desde 2016 terminó por encima del 30%.", [("Por eso conviene no subir el tope de cuota.", {"bold": True, "color": INK, "espacio": 12})]], size=18, color=INK2)
    pie(s, "Fuente: BCRA (UVA), INDEC (índice de salarios registrados); elaboración propia.")

    s = d.slide()
    d.texto(s, m, 0.8, w, 0.5, ["RECOMENDACIÓN"], size=13, bold=True, color=BLUE)
    d.texto(s, m, 1.35, w, 3.6, ["Crédito a 30 años, 80% financiado y cuota hasta 25%, con foco en hogares con dos ingresos y una línea para independientes."], size=28, bold=True, interlineado=1.05)
    d.texto(s, m, 5.2, w, 2.6, [f"•  De {r['elegibles']} a {r['prop_indep']} hogares elegibles", "•  Baja el efectivo inicial que hay que juntar", "•  Mantiene el tope de cuota frente al riesgo de desacople"], size=18, color=INK2, interlineado=1.4)
    pie(s, "Caso con banco ficticio. No es asesoramiento financiero.")

    s = d.slide()
    d.texto(s, m, 1.2, w, 1.6, ["Recorré el caso completo"], size=36, bold=True)
    d.texto(s, m, 2.9, w, 3.4, ["•  Dashboard con simulador", "•  Resumen en simple", "•  Código, documentación y fuentes en GitHub", "•  Decisiones y controles en cada etapa"], size=20, color=INK2, interlineado=1.5)
    d.texto(s, m, 7.4, w, 0.9, [[(AUTOR, {"bold": True, "color": INK}), ("  ·  link en el post", {"color": MUTED})]], size=16)
    return d


def exportar_pdf(pptx: Path, titulo: str) -> Path | None:
    if not (SOFFICE.exists() or shutil.which(str(SOFFICE))):
        print("  (LibreOffice no encontrado: se omite el PDF)")
        return None
    subprocess.run([str(SOFFICE), "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
                   check=True, timeout=300, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = pptx.with_suffix(".pdf")
    limpiar_metadatos(pdf, titulo)
    return pdf


def main() -> None:
    K = json.loads((PROCESSED / "resultados" / "kpis.json").read_text(encoding="utf-8"))
    r, prot = resultados(K), protagonista(K)
    DECK.mkdir(exist_ok=True)
    for nombre, constructor, titulo in [("presentacion", presentacion, "Presentación · Crédito hipotecario UVA"),
                                        ("carrusel_linkedin", carrusel, "Carrusel · Crédito hipotecario UVA")]:
        pptx = DECK / f"{nombre}.pptx"
        constructor(r, prot, K).guardar(pptx)
        pdf = exportar_pdf(pptx, titulo)
        print(f"  ✓ deck/{pptx.name}" + (f" + {pdf.name} ({pdf.stat().st_size / 1e6:.1f} MB)" if pdf else ""))


if __name__ == "__main__":
    main()
